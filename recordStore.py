"""Saves weighing records to one Excel file per day."""
import csv
import logging
import os
import threading
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

log = logging.getLogger(__name__)

SHEET_NAME = "Weighings"
JOURNAL_FIELDS = ["id", "timestamp", "scale", "weight"]
EXTRA_COLUMNS = ["Seller", "Material", "Notes"]  # left blank for the user to fill in
COLUMN_WIDTHS = [26, 12, 10, 14, 13, 22, 22, 32]


class RecordStore:
    """Saves each weighing to the day's Excel workbook, without ever losing one.

    Every record is first appended to a CSV journal (fast and crash-safe), then
    copied into the workbook. If the workbook can't be written, e.g. because
    it's open in Excel, the record stays in the journal and is copied on the
    next sync. Rows are only ever appended, so anything typed into the Seller,
    Material and Notes columns is kept.

    Methods are safe to call from several scale threads at once.
    """

    def __init__(self, output_dir, unit="kg"):
        self.output_dir = Path(output_dir)
        self.journal_dir = self.output_dir / "journal"
        self.unit = unit
        self._lock = threading.Lock()
        self._pending_days = set()
        self._problem_files = set()  # workbooks already warned about, to avoid repeating it

    def workbook_path(self, day):
        return self.output_dir / f"scale_data_{day:%Y-%m-%d}.xlsx"

    def _journal_path(self, day):
        return self.journal_dir / f"{day:%Y-%m-%d}.csv"

    def add(self, scale, weight, timestamp):
        day = timestamp.date()
        with self._lock:
            record = {
                # Numbered per day, like a ticket number: 20260924-0001, 20260924-0002, ...
                "id": f"{day:%Y%m%d}-{self._journal_count(day) + 1:04d}",
                "timestamp": timestamp.isoformat(timespec="seconds"),
                "scale": scale,
                "weight": weight,
            }
            self._append_to_journal(record, day)
            self._pending_days.add(day)
            self._sync_pending()
        return record

    def sync(self):
        """Retries copying records into workbooks that couldn't be written before."""
        with self._lock:
            self._sync_pending()

    def sync_recent(self, days=31):
        """Checks the last few days' workbooks against their journals.

        Call at startup to recover records that were journaled but never reached
        Excel, e.g. because the program was closed while a workbook was open.
        """
        cutoff = date.today() - timedelta(days=days)
        with self._lock:
            for journal in self.journal_dir.glob("*.csv"):
                try:
                    day = date.fromisoformat(journal.stem)
                except ValueError:
                    continue
                if day >= cutoff:
                    self._pending_days.add(day)
            self._sync_pending()

    def _journal_count(self, day):
        path = self._journal_path(day)
        if not path.exists():
            return 0
        with open(path, newline="", encoding="utf-8") as f:
            return sum(1 for _ in csv.DictReader(f))

    def _append_to_journal(self, record, day):
        self.journal_dir.mkdir(parents=True, exist_ok=True)
        path = self._journal_path(day)
        is_new = not path.exists()
        with open(path, "a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, JOURNAL_FIELDS)
            if is_new:
                writer.writeheader()
            writer.writerow(record)
            f.flush()
            os.fsync(f.fileno())

    def _sync_pending(self):
        for day in sorted(self._pending_days):
            if self._sync_day(day):
                self._pending_days.discard(day)

    def _sync_day(self, day):
        """Adds the day's journaled records missing from its workbook.

        Returns True once the workbook contains every record.
        """
        journal = self._journal_path(day)
        if not journal.exists():
            return True
        with open(journal, newline="", encoding="utf-8") as f:
            records = list(csv.DictReader(f))

        path = self.workbook_path(day)
        try:
            workbook = load_workbook(path) if path.exists() else self._new_workbook()
            sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
            saved_ids = {str(row[0]) for row in
                         sheet.iter_rows(min_row=2, max_col=1, values_only=True)}
            missing = [r for r in records if r["id"] not in saved_ids]
            if not missing:
                return True
            for record in missing:
                self._append_row(sheet, record)
            # Write to a temporary file first so a crash mid-save can't corrupt the workbook.
            temp = path.with_name(f"~{path.name}")
            workbook.save(temp)
            os.replace(temp, path)
        except PermissionError:
            if path not in self._problem_files:
                log.warning("%s is open in another program; new records will be added "
                            "to it once it's closed", path.name)
                self._problem_files.add(path)
            return False
        except Exception:
            if path not in self._problem_files:
                log.exception("Could not update %s; records are kept in %s and will be "
                              "retried", path.name, journal)
                self._problem_files.add(path)
            return False

        if path in self._problem_files:
            log.info("%s is up to date again", path.name)
            self._problem_files.discard(path)
        return True

    def _new_workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(["Record ID", "Date", "Time", "Scale", f"Weight ({self.unit})",
                      *EXTRA_COLUMNS])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for cell, width in zip(sheet[1], COLUMN_WIDTHS):
            sheet.column_dimensions[cell.column_letter].width = width
        sheet.freeze_panes = "A2"
        return workbook

    def _append_row(self, sheet, record):
        timestamp = datetime.fromisoformat(record["timestamp"])
        sheet.append([record["id"], timestamp.date(), timestamp.time(), record["scale"],
                      float(record["weight"])])
        row = sheet.max_row
        sheet.cell(row, 2).number_format = "yyyy-mm-dd"
        sheet.cell(row, 3).number_format = "hh:mm:ss"
