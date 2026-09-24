import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from recordStore import RecordStore

MORNING = datetime(2026, 9, 24, 9, 30, 0)
NOON = datetime(2026, 9, 24, 12, 0, 0)


def read_rows(path):
    sheet = load_workbook(path).active
    return [row for row in sheet.iter_rows(min_row=2, values_only=True)]


class RecordStoreTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir = Path(self.temp_dir.name)
        self.store = RecordStore(self.dir)
        self.workbook = self.store.workbook_path(MORNING.date())

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_records_are_appended_to_the_days_workbook(self):
        self.store.add("WT1000", 12.5, MORNING)
        self.store.add("WT3000i", 300.0, NOON)
        rows = read_rows(self.workbook)
        self.assertEqual([(r[3], r[4]) for r in rows], [("WT1000", 12.5), ("WT3000i", 300.0)])
        self.assertEqual(self.workbook.name, "scale_data_2026-09-24.xlsx")

    def test_records_in_the_same_second_are_all_kept(self):
        self.store.add("WT1000", 12.5, MORNING)
        self.store.add("WT1000", 13.0, MORNING)
        rows = read_rows(self.workbook)
        self.assertEqual([(r[0], r[4]) for r in rows],
                         [("20260924-0001", 12.5), ("20260924-0002", 13.0)])

    def test_details_typed_by_the_user_are_kept(self):
        self.store.add("WT1000", 12.5, MORNING)
        workbook = load_workbook(self.workbook)
        workbook.active["F2"] = "Seller A"
        workbook.save(self.workbook)

        self.store.add("WT1000", 20.0, NOON)
        rows = read_rows(self.workbook)
        self.assertEqual(rows[0][5], "Seller A")
        self.assertEqual(len(rows), 2)

    def test_records_wait_while_the_workbook_is_open(self):
        self.store.add("WT1000", 12.5, MORNING)
        with mock.patch("recordStore.os.replace", side_effect=PermissionError):
            self.store.add("WT1000", 20.0, NOON)
        self.assertEqual(len(read_rows(self.workbook)), 1)

        self.store.sync()
        self.assertEqual([r[4] for r in read_rows(self.workbook)], [12.5, 20.0])

    def test_startup_recovers_records_missing_from_excel(self):
        self.store.add("WT1000", 12.5, MORNING)
        self.workbook.unlink()

        with mock.patch("recordStore.date") as fake_date:
            fake_date.today.return_value = MORNING.date()
            fake_date.fromisoformat = MORNING.date().fromisoformat
            RecordStore(self.dir).sync_recent()
        self.assertEqual([r[4] for r in read_rows(self.workbook)], [12.5])


if __name__ == "__main__":
    unittest.main()
